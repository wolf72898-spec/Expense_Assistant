"""生成《远筑工程 · 工程变更与签证协同 AI 试点 报价单》xlsx。

金额在脚本内按「人天 × 单价」计算后写入（不用公式），
保证在任何表格软件里打开都能直接看到数字，不依赖重算。

用法：
    python scripts/gen_quote_xlsx.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "fde-交付物" / "06-报价单-工程变更与签证协同AI试点.xlsx"

BLUE = "1C5FB0"
LIGHT = "EAF2FB"
GREY_F = "F5F7FA"
TOTAL_F = "FFF6E5"

thin = Side(style="thin", color="C9D3E0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# 角色单价（元 / 人天）
RATE = {
    "解决方案架构师": 3500,
    "产品经理": 2500,
    "AI / 后端工程师": 2800,
    "数据 / 集成工程师": 2500,
    "内核迁移工程师": 2800,
    "测试工程师": 1800,
    "项目经理": 2500,
}

# (阶段, 工作项, 角色, 人天)
ITEMS = [
    ("P0 诊断与基线", "业务诊断与访谈复盘", "解决方案架构师", 6),
    ("P0 诊断与基线", "六节点耗时基线计量设计", "产品经理", 8),
    ("P0 诊断与基线", "字段与主编号盘点", "数据 / 集成工程师", 6),
    ("P0 诊断与基线", "项目管理", "项目经理", 3),

    ("P1 只读 MVP", "内核迁移：规则中枢（版本化 + 审批分离 + 变更审计）", "内核迁移工程师", 10),
    ("P1 只读 MVP", "内核迁移：溯源标签与硬阻断闸门", "内核迁移工程师", 8),
    ("P1 只读 MVP", "M1 版本核对与交底提示智能体", "AI / 后端工程师", 18),
    ("P1 只读 MVP", "M2 依据与授权链预检智能体", "AI / 后端工程师", 16),
    ("P1 只读 MVP", "M7 人工修复队列与歧义处理", "数据 / 集成工程师", 12),
    ("P1 只读 MVP", "现场终端草稿暂存与最后同步时间", "数据 / 集成工程师", 8),
    ("P1 只读 MVP", "产品与交互设计", "产品经理", 10),
    ("P1 只读 MVP", "测试与回放验证", "测试工程师", 10),
    ("P1 只读 MVP", "项目管理", "项目经理", 6),

    ("P2 阻断与集成", "M4 阻断提交与人工转办", "AI / 后端工程师", 14),
    ("P2 阻断与集成", "M5 跨系统只读集成（BIM 库 / 项目平台 / 造价软件）", "数据 / 集成工程师", 20),
    ("P2 阻断与集成", "M3 计量口径归一与差异定位", "AI / 后端工程师", 12),
    ("P2 阻断与集成", "快照时效与故障降级规则", "数据 / 集成工程师", 8),
    ("P2 阻断与集成", "产品与交互设计", "产品经理", 8),
    ("P2 阻断与集成", "测试", "测试工程师", 12),
    ("P2 阻断与集成", "项目管理", "项目经理", 6),

    ("P3 试点与验收", "M6 留痕归档与审计视图", "AI / 后端工程师", 12),
    ("P3 试点与验收", "历史资料回放与试点验证", "测试工程师", 8),
    ("P3 试点与验收", "分区试点支持与验收", "产品经理", 8),
    ("P3 试点与验收", "运维手册与运营交接", "解决方案架构师", 6),
    ("P3 试点与验收", "项目管理", "项目经理", 8),
]

OTHERS = [
    ("云资源与大模型调用（18 周 · 单分区估算）", 45000),
    ("差旅与现场支持", 30000),
    ("培训与宣贯", 15000),
]

DISCOUNT_RATE = 0.10   # 首个分区试点优惠
TAX_RATE = 0.06        # 增值税


def style(ws, cell_range, *, bold=False, size=10.5, fill=None, color=None,
          align="left", wrap=True, numfmt=None):
    for row in ws[cell_range]:
        for c in row:
            c.font = Font(name="微软雅黑", size=size, bold=bold,
                          color=color or "1F2A37")
            c.border = BORDER
            c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
            if numfmt:
                c.number_format = numfmt


def title_block(ws):
    ws["A1"] = "远筑工程 · 工程变更与签证协同 AI 试点"
    ws["A1"].font = Font(name="微软雅黑", size=16, bold=True, color=BLUE)
    ws["A2"] = "项目报价单"
    ws["A2"].font = Font(name="微软雅黑", size=13, bold=True, color=BLUE)
    meta = [
        ("报价单号", "QH-FDE-2026-001"),
        ("报价日期", "2026-10-30"),
        ("有效期", "自报价日起 30 日"),
        ("币种", "人民币（元）"),
        ("试点范围", "一个施工分区（历史资料回放 + 并行试点）"),
    ]
    r = 4
    for k, v in meta:
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
        r += 1
    style(ws, "A4:B8", fill=GREY_F, size=10)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 16


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"
    title_block(ws)

    # 表头
    head_row = 10
    headers = ["序号", "工作项", "角色", "人天", "单价（元/人天）", "小计（元）"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=head_row, column=i, value=h)
    style(ws, f"A{head_row}:F{head_row}", bold=True, fill=LIGHT, align="center")

    row = head_row + 1
    idx = 0
    stage_start = {}
    labor = 0
    days = 0
    for stage, item, role, d in ITEMS:
        stage_start.setdefault(stage, row)
        idx += 1
        rate = RATE[role]
        sub = d * rate
        labor += sub
        days += d
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=f"【{stage}】{item}")
        ws.cell(row=row, column=3, value=role)
        ws.cell(row=row, column=4, value=d)
        ws.cell(row=row, column=5, value=rate)
        ws.cell(row=row, column=6, value=sub)
        row += 1
    style(ws, f"A{head_row+1}:F{row-1}", size=10)
    style(ws, f"D{head_row+1}:D{row-1}", align="center", wrap=False)
    style(ws, f"E{head_row+1}:F{row-1}", align="right", wrap=False, numfmt="#,##0")

    # 人力小计
    ws.cell(row=row, column=2, value="人力费用小计")
    ws.cell(row=row, column=4, value=days)
    ws.cell(row=row, column=6, value=labor)
    style(ws, f"A{row}:F{row}", bold=True, fill=LIGHT)
    style(ws, f"D{row}:D{row}", align="center", wrap=False)
    style(ws, f"F{row}:F{row}", align="right", wrap=False, numfmt="#,##0")
    labor_row = row
    row += 2

    # 非人力
    ws.cell(row=row, column=2, value="非人力费用")
    style(ws, f"A{row}:F{row}", bold=True, fill=LIGHT)
    row += 1
    other_total = 0
    for name, amt in OTHERS:
        other_total += amt
        ws.cell(row=row, column=2, value=name)
        ws.cell(row=row, column=6, value=amt)
        style(ws, f"A{row}:F{row}", size=10)
        style(ws, f"F{row}:F{row}", align="right", wrap=False, numfmt="#,##0")
        row += 1
    ws.cell(row=row, column=2, value="非人力费用小计")
    ws.cell(row=row, column=6, value=other_total)
    style(ws, f"A{row}:F{row}", bold=True, fill=LIGHT)
    style(ws, f"F{row}:F{row}", align="right", wrap=False, numfmt="#,##0")
    row += 2

    # 汇总
    subtotal = labor + other_total
    discount = round(subtotal * DISCOUNT_RATE)
    taxed_base = subtotal - discount
    tax = round(taxed_base * TAX_RATE)
    total = taxed_base + tax

    ws.cell(row=row, column=2, value="费用合计（不含税）")
    ws.cell(row=row, column=6, value=subtotal)
    style(ws, f"A{row}:F{row}", align="right", wrap=False, numfmt="#,##0")
    row += 1
    ws.cell(row=row, column=2, value=f"首个分区试点优惠（-{int(DISCOUNT_RATE*100)}%）")
    ws.cell(row=row, column=6, value=-discount)
    style(ws, f"A{row}:F{row}", numfmt="#,##0")
    style(ws, f"F{row}:F{row}", align="right", wrap=False)
    row += 1
    ws.cell(row=row, column=2, value=f"增值税（{int(TAX_RATE*100)}%）")
    ws.cell(row=row, column=6, value=tax)
    style(ws, f"A{row}:F{row}", numfmt="#,##0")
    style(ws, f"F{row}:F{row}", align="right", wrap=False)
    row += 1
    ws.cell(row=row, column=2, value="报价总计")
    ws.cell(row=row, column=6, value=total)
    style(ws, f"A{row}:F{row}", bold=True, size=12, fill=TOTAL_F, color=BLUE)
    style(ws, f"F{row}:F{row}", align="right", wrap=False, numfmt="#,##0")
    row += 2

    # 付款节奏
    ws.cell(row=row, column=2, value="付款节奏")
    style(ws, f"A{row}:F{row}", bold=True, fill=LIGHT)
    row += 1
    for name, ratio in [("合同签订后 5 个工作日内", 0.30),
                        ("P1 只读 MVP 验收通过后", 0.40),
                        ("P3 试点验收与交接完成后", 0.30)]:
        ws.cell(row=row, column=2, value=name)
        ws.cell(row=row, column=3, value=f"{int(ratio*100)}%")
        ws.cell(row=row, column=6, value=round(total * ratio))
        style(ws, f"A{row}:F{row}", size=10)
        style(ws, f"C{row}:C{row}", align="center", wrap=False)
        style(ws, f"F{row}:F{row}", align="right", wrap=False, numfmt="#,##0")
        row += 1
    row += 1

    # 不含项
    ws.cell(row=row, column=2, value="本报价不含以下项目")
    style(ws, f"A{row}:F{row}", bold=True, fill=LIGHT)
    row += 1
    for s in [
        "BIM 建模、模型重建与工程量重新计算服务",
        "第三方系统 License 费用，以及 BIM 文档库 / 项目管理平台 / 造价软件的接口改造费用",
        "现场硬件、网络与终端采购",
        "数据脱敏规则与留痕保留期限的制定（由甲方制度部门确定，平台只执行）",
        "质量与安全专业判定服务（AI 建议不得作为判定依据，须由具备资格的人员签核）",
    ]:
        ws.cell(row=row, column=2, value=f"· {s}")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        style(ws, f"A{row}:F{row}", size=10)
        row += 1
    row += 1

    # 假设条件
    ws.cell(row=row, column=2, value="报价假设条件")
    style(ws, f"A{row}:F{row}", bold=True, fill=LIGHT)
    row += 1
    for s in [
        "试点范围为一个施工分区，先做历史资料回放，再做小范围并行试点，不直接替换现有业务",
        "甲方指定接口人与业务负责人，评审反馈不超过 3 个工作日，否则工期顺延",
        "三套系统需提供只读接口或导出能力；如需接口改造，另行评估报价",
        "大模型调用量按 18 周、单分区规模估算，超出部分按实际用量结算",
        "各节点耗时基线由双方在 P0 阶段共同采集；基线缺失将导致 6 天目标无法验收",
        "本报价基于合成场景访谈所得需求假设，正式合同前须完成现场核实与基线确认",
    ]:
        ws.cell(row=row, column=2, value=f"· {s}")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        style(ws, f"A{row}:F{row}", size=10)
        row += 1

    # 单价表
    ws2 = wb.create_sheet("角色单价表")
    ws2["A1"] = "角色"
    ws2["B1"] = "单价（元 / 人天）"
    style(ws2, "A1:B1", bold=True, fill=LIGHT, align="center")
    r = 2
    for k, v in RATE.items():
        ws2.cell(row=r, column=1, value=k)
        ws2.cell(row=r, column=2, value=v)
        style(ws2, f"A{r}:B{r}", size=10.5)
        style(ws2, f"B{r}:B{r}", align="right", wrap=False, numfmt="#,##0")
        r += 1
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 18

    ws.freeze_panes = "A11"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(OUT))

    print(f"已生成：{OUT}")
    print(f"  人力 {days} 人天 / ¥{labor:,}")
    print(f"  非人力 ¥{other_total:,}")
    print(f"  合计 ¥{subtotal:,} → 优惠 -¥{discount:,} → 税 ¥{tax:,}")
    print(f"  报价总计 ¥{total:,}")


if __name__ == "__main__":
    build()
